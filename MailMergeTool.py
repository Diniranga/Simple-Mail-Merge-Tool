from docx import Document
import csv
import copy
import re
import os
import sys

def replace_placeholders(doc, placeholders, data):
    name = ''
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            text = run.text
            if '[' in text and ']' in text:
                matches = re.findall(r'\[(\d+)\]', text)
                for match in matches:
                    placeholder = f'[{match}]'
                    if placeholder in placeholders:
                        index = placeholders.index(placeholder)
                        value = str(data[index])

                        if index == 1:
                            name = value
                        if index == 2:
                            address_parts = value.split(',')
                            run.text = '\n'.join(address_parts)
                        else:
                            run.text = run.text.replace(placeholder, value)
    return name

def main():
    script_folder = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(os.path.abspath(__file__))
    word_file_path = os.path.join(script_folder, 'template.docx')
    csv_file_path = os.path.join(script_folder, 'data.csv')
    output_folder = os.path.join(script_folder, 'output')

    with open(csv_file_path, 'r') as csv_file:
        csv_reader = csv.reader(csv_file)
        data_rows = list(csv_reader)

    placeholders = data_rows[0]
    doc = Document(word_file_path)

    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    count = 0
    for data in data_rows[1:]:
        current_doc = copy.deepcopy(doc)
        employee = replace_placeholders(current_doc, placeholders, data)
        output_word_file_path = os.path.join(output_folder, f'{employee}.docx')
        current_doc.save(output_word_file_path)
        print(f"Word file with replaced data for row {count + 1} saved to {output_word_file_path}")
        count += 1

if __name__ == "__main__":
    main()
